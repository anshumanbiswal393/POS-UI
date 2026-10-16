"""Generate Microsoft Word plan for POS backend inside PMS."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT_PATHS = [
    Path(r"D:\Retrod_Work\POS\docs\Retrod_POS_Backend_PMS_Integration_Plan.docx"),
    Path(
        r"D:\Retrod_Work\PMS\backend\retrod-one-backend\apps\POS"
        r"\Retrod_POS_Backend_PMS_Integration_Plan.docx"
    ),
]


def add_p(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = "Calibri"
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Calibri"
    doc.add_paragraph()


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Retrod POS Backend")
    r.bold = True
    r.font.size = Pt(22)
    r.font.name = "Calibri"
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Part-wise Build Plan inside PMS Backend\n"
        "Same architecture as PMS · Room Charge → Guest Folio"
    )
    r.font.size = Pt(12)
    r.font.name = "Calibri"

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Document Version 1.0  |  Date: 16 Oct 2026  |  "
        "Status: Plan only — coding starts after review"
    )
    r.font.size = Pt(10)
    r.italic = True

    doc.add_paragraph()

    doc.add_heading("1. Brief summary of all work", level=1)
    add_p(
        doc,
        "We will build the Retrod POS backend as a Django app inside the existing "
        "PMS backend (not a separate server). POS UI will call /api/pos/... APIs. "
        "PMS remains the source of truth for rooms, check-in, and guest folio. "
        "POS owns restaurants, menus, orders, KOT, and payments. When a guest "
        "chooses Room Charge, POS posts the food bill onto the open Guest Folio "
        "of the currently occupied room so Front Office sees it at checkout.",
    )
    add_p(doc, "Work in short (what we will deliver):", bold=True)
    add_numbered(
        doc,
        [
            "Harden apps/POS foundation using the same code layout PMS feature apps use.",
            "Occupied-rooms bridge — fetch in-house rooms/guests/folios from PMS for Room Service.",
            "Menu catalog APIs for Billing / Menu screens.",
            "Orders + KOT + Current Orders / Order History APIs aligned with POS UI flow.",
            "Payments including Room Charge that posts to GuestFolio / FolioTransaction.",
            "Due bills, exports, then later stock/recipes/reports.",
        ],
    )

    doc.add_heading("2. Where the POS backend lives", level=1)
    add_table(
        doc,
        ["Item", "Location"],
        [
            ["PMS backend root", "PMS/backend/retrod-one-backend/"],
            ["POS Django app", "apps/POS/  (INSTALLED_APPS: apps.POS, label: pos)"],
            ["API mount", "/api/pos/  via retrod_pms/urls.py"],
            ["POS UI", "POS/ (Vite + React)"],
            [
                "Already built",
                "Outlet, FloorSection, RestaurantTable, POSRegister, DeviceMapping + CRUD",
            ],
        ],
    )

    doc.add_heading("3. Follow the same backend architecture as PMS", level=1)
    add_p(
        doc,
        "Every new POS module must mirror how PMS feature apps are structured today "
        "(example: apps/features/front_office/). We will not invent a different pattern.",
    )
    add_p(doc, "PMS feature app pattern (to copy):", bold=True)
    add_bullets(
        doc,
        [
            "apps.py — AppConfig",
            "models.py — Django models on BaseModel (tenant, soft delete, timestamps)",
            "serializers.py — DRF serializers",
            "views.py — ViewSets / APIViews + drf-spectacular tags",
            "services.py (or services/) — business logic & cross-app calls (not fat views)",
            "urls.py — router registration",
            "tests.py / tests/ — API + service tests",
            "migrations/ — only app-owned migrations",
        ],
    )
    add_p(doc, "Target layout for apps/POS (same spirit as PMS):", bold=True)
    add_bullets(
        doc,
        [
            "apps/POS/models.py — POS-owned tables only",
            "apps/POS/serializers.py",
            "apps/POS/views.py",
            "apps/POS/urls.py — included at /api/pos/",
            "apps/POS/services/occupancy.py — read CHECKED_IN rooms from PMS",
            "apps/POS/services/folio_bridge.py — Room Charge post to GuestFolio",
            "apps/POS/permissions.py — RBAC codes (pos:*)",
            "apps/POS/tests/ — part-wise tests",
        ],
    )
    add_p(doc, "Shared PMS building blocks POS must reuse:", bold=True)
    add_bullets(
        doc,
        [
            "BaseModel from apps.core.common.models",
            "Tenant / Property from apps.core.tenants",
            "JWT auth + RBAC from apps.core.accounts / apps.core.rbac",
            "Reservations (CHECKED_IN) from apps.features.reservations",
            "Rooms from apps.features.inventory",
            "GuestFolio + FolioTransaction from apps.features.front_office",
            "Same headers: X-Tenant-ID, X-Property-ID",
            "Same OpenAPI style via drf-spectacular tags (e.g. POS - Orders)",
        ],
    )

    doc.add_heading("4. Ownership rules (no conflict with PMS)", level=1)
    add_table(
        doc,
        ["Domain", "Owner", "POS role"],
        [
            ["Auth / Tenant / RBAC", "apps.core.*", "Reuse only"],
            ["Rooms / Housekeeping", "inventory + housekeeping", "Read occupied status"],
            ["Stay / Check-in", "reservations + front_office", "Read CHECKED_IN stays"],
            ["Guest Folio balance", "front_office", "Post via FolioBridge service only"],
            ["Menu / Tables / KOT / Orders", "apps.POS", "Full owner"],
            ["Night audit / Checkout", "front_office", "Do not bypass"],
        ],
    )
    add_p(doc, "Hard rules:", bold=True)
    add_numbered(
        doc,
        [
            "Do not create a second guest/stay/folio model inside POS.",
            "Do not edit old PMS migrations — only makemigrations pos.",
            "Foreign keys are one-way: POS → Reservation / Unit / Folio.",
            "Room Charge settle must use transaction.atomic() + idempotency key.",
            "Views stay thin; business logic goes in services.py (PMS style).",
            "Circular imports avoided: POS services call front_office services, not views.",
        ],
    )

    doc.add_heading("5. Room Service → Room Charge → Guest Folio flow", level=1)
    add_numbered(
        doc,
        [
            "Guest is checked in (reservation CHECKED_IN) with an OPEN GuestFolio.",
            "POS UI loads occupied rooms from GET /api/pos/integrations/occupied-rooms/ (reads PMS data).",
            "Staff takes Room Service / restaurant order → POST /api/pos/orders/.",
            "KOT sent to kitchen → order stays on Current Orders in Preparation.",
            "Guest asks to put bill on room → settle with payment_mode = ROOM_CHARGE.",
            "FolioBridge validates stay + open folio, creates FolioTransaction (CHARGE), "
            "updates folio balance, marks POS order SETTLED, stores folio_transaction_id.",
            "Front Office guest folio shows the food line; checkout includes it automatically.",
        ],
    )
    add_p(doc, "Failure safeguards:", bold=True)
    add_bullets(
        doc,
        [
            "Reject Room Charge if folio is SETTLED/SUSPENDED or guest already checked out.",
            "Same idempotency key → no double posting.",
            "Void after room charge → REVERSAL FolioTransaction + manager permission.",
        ],
    )

    doc.add_heading("6. Part-wise backend work (do in this order)", level=1)

    doc.add_heading("Part 0 — Foundation (align with PMS architecture)", level=2)
    add_bullets(
        doc,
        [
            "Confirm apps.POS registration and /api/pos/ mount (already done).",
            "Add services/, permissions.py, tests/ using PMS feature-app pattern.",
            "Keep AllowAny out of production paths; use RBAC.",
            "Exit: manage.py check green; existing outlet/table APIs still work.",
        ],
    )

    doc.add_heading("Part 1 — Occupied rooms bridge", level=2)
    add_bullets(
        doc,
        [
            "Service: services/occupancy.py reading CHECKED_IN reservations + room + guest + open folio.",
            "API: GET /api/pos/integrations/occupied-rooms/",
            "Wire POS Room Service / billing room picker.",
            "Exit: In-house PMS room appears in POS; vacant rooms do not.",
        ],
    )

    doc.add_heading("Part 2 — Menu catalog", level=2)
    add_bullets(
        doc,
        [
            "POS-owned models: categories, items, variants, modifiers, tax profile.",
            "APIs under /api/pos/menu/...",
            "Seed command; wire Billing catalog UI.",
        ],
    )

    doc.add_heading("Part 3 — Orders + KOT + Current Orders", level=2)
    add_bullets(
        doc,
        [
            "Models: POSOrder, POSOrderItem, KOTTicket, status history.",
            "Flow: Placed → Prepared → Settled (matches POS UI).",
            "Edit open order → incremental KOT + updated totals.",
            "APIs for current list, history, advance stage, settle.",
        ],
    )

    doc.add_heading("Part 4 — Payments + Room Charge to folio (critical)", level=2)
    add_bullets(
        doc,
        [
            "POSPayment + FolioPostingLog (unique idempotency_key).",
            "services/folio_bridge.py posts to GuestFolio / FolioTransaction.",
            "Charge codes: POS_ROOM_SERVICE / POS_FNB (agree once with FO).",
            "Exit tests: one folio line, balance up, no double post on retry.",
        ],
    )

    doc.add_heading("Part 5 — Due bills, splits, day close, export", level=2)
    add_bullets(
        doc,
        [
            "Pending payments ledger, split tender, day-end totals, Excel export.",
        ],
    )

    doc.add_heading("Part 6 — Stock / recipes / advanced reports (later)", level=2)
    add_bullets(
        doc,
        [
            "BOM deductions and reports after core order + folio path is stable.",
        ],
    )

    doc.add_heading("7. Backend part ↔ POS UI screen", level=1)
    add_table(
        doc,
        ["Order", "Backend part", "POS UI"],
        [
            ["1", "Occupied rooms", "Room Service, Billing room picker"],
            ["2", "Menu APIs", "Billing catalog, Menu On/Off"],
            ["3", "Orders + KOT", "Take Order, Billing, Current Orders, KDS"],
            ["4", "Room Charge folio", "Settle → Room Charge; PMS Guest Folio"],
            ["5", "History / dues / export", "Order History, Pending Payments"],
            ["6", "Stock & reports", "Inventory, Reports"],
        ],
    )

    doc.add_heading("8. Key API sketches", level=1)
    add_p(doc, "Occupied rooms:", bold=True)
    add_p(doc, "GET /api/pos/integrations/occupied-rooms/?property_id=...")
    add_p(doc, "Settle with room charge:", bold=True)
    add_p(
        doc,
        'POST /api/pos/orders/{id}/settle/  body: '
        '{"payment_mode":"ROOM_CHARGE","reservation_id":"...","folio_id":"...",'
        '"idempotency_key":"..."}',
    )

    doc.add_heading("9. How we will work (process)", level=1)
    add_numbered(
        doc,
        [
            "Review this Word plan — then start Part 1 only.",
            "Branch from latest PMS backend; one PR per part.",
            "Follow PMS file layout: models → serializers → services → views → urls → tests.",
            "makemigrations pos only; never rewrite PMS migrations.",
            "Do not mix PMS reservation refactors with POS order features in the same PR.",
            "Demo each part: PMS check-in → POS room list → order → Room Charge → folio line in FO.",
        ],
    )

    doc.add_heading("10. Next step after this document", level=1)
    add_p(
        doc,
        "No coding until this plan is accepted. First implementation task will be Part 1: "
        "Occupied rooms bridge (services/occupancy.py + GET /api/pos/integrations/occupied-rooms/), "
        "written in the same architectural style as PMS front_office / reservations apps.",
    )

    doc.add_heading("11. Related files", level=1)
    add_bullets(
        doc,
        [
            "Markdown twin: POS/docs/POS_BACKEND_PMS_INTEGRATION_PLAN.md",
            "POS UI roadmap: POS/docs/POS_IMPLEMENTATION_ROADMAP.md",
            "Existing POS seed: apps/POS/models.py, urls.py, views.py",
            "Folio models: apps/features/front_office/models.py (GuestFolio, FolioTransaction)",
        ],
    )

    end = doc.add_paragraph()
    run = end.add_run("— End of plan —")
    run.italic = True

    for path in OUT_PATHS:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        print(f"Wrote: {path}")


if __name__ == "__main__":
    build()
